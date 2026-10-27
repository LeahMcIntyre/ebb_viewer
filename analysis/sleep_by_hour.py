"""Sleep by hour: turn a finished ebb-viewer project into sleep tables, Prism workbooks and plots.

This is the script version of ``sleep_by_hour.ipynb`` (same calculation, same output).
The step-by-step explanation of the calculation is in ``sleep_by_hour_walkthrough.ipynb``.

HOW TO RUN
    1. Make a copy of this file for each analysis.
    2. Edit the block marked "SETTINGS" below. Lines marked  <<< CHANGE THIS  must be edited;
       everything else has a sensible default.
    3. Run it:      python sleep_by_hour.py
       The script checks your settings first and stops with a plain message if one is wrong.

WHAT YOU GET, in OUT_DIR (default: <project>/sleep_analysis/)
    prism_sleep_by_hour.xlsx   one sheet per state x metric, rows = ZT bins, animals blocked by group
    prism_sleep_whole_file.xlsx   one number per animal and state for the whole recording
    sleep_bins.csv, sleep_bins_avg.csv, sleep_whole_file.csv   the same data as tidy tables
    sleep_by_hour.png, sleep_whole_file.png   the group plots

HOW THE NUMBERS ARE MADE
    The reclassified state (``final_state``) of each 4 s epoch is repeated to one state per second,
    the recording is rotated so its first 6 AM is position 0, and it is cut into 2-hour bins. Per
    bin and state: ``frac_time`` (fraction of the bin), ``n_bouts`` (unbroken runs) and
    ``mean_bout_sec``. Bins are folded onto a 24 h day and averaged across days;
    ZT = (clock hour - 6) mod 24. Animals excluded in Sleep QC are left out.
"""

import csv
import json
import sys
from datetime import datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import matplotlib

matplotlib.use("Agg")  # write the plots to files; never open a window
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook

# =====================================================================================
# SETTINGS  -  EDIT THIS BLOCK. Everything below the next line of "=" is the calculation.
# =====================================================================================

# <<< CHANGE THIS: the project folder, i.e. the folder that holds manifest.json.
PROJECT_ROOT = Path("/path/to/your/project")

# <<< CHANGE THIS (or leave both empty for one group with everyone): the experimental groups.
# Either list the animals here ...
#     {"group name": [animal IDs]}
GROUPS = {
    # "WT":  ["Mice-1248", "Mice-1253"],
    # "MUT": ["Mice-1251", "Mice-1279"],
}
# ... or read them from a CSV file (this replaces GROUPS). The easiest layout is one row per
# animal, with two columns named "animal" and "group":
#     animal,group
#     Mice-1248,WT
#     Mice-1251,MUT
# One column per group also works: the group name on top, its animal IDs below it:
#     WT,MUT
#     Mice-1248,Mice-1251
#     Mice-1253,Mice-1279
GROUPS_CSV = None  # e.g. Path("/path/to/groups.csv")

# <<< CHECK THIS: the time zone the recordings were made in (used for daylight saving).
# None = ignore daylight saving.
TZ = "America/Chicago"

# ----- optional; the defaults are what the lab normally uses -----
OUT_DIR = None  # None -> <PROJECT_ROOT>/sleep_analysis. Or Path("/path/to/results")
INCLUDE_EXCLUDED = False  # True keeps animals excluded in Sleep QC (marked in an "excluded" column)
BIN_HOURS = 2  # width of each bin; must divide 24
BEGIN_HR = 6  # clock hour of ZT 0 (6 = lights on at 6 AM)
DARK_ONSET_ZT = 12  # lights-off, for the grey band on the plots
STATE_COLUMN = "final_state"  # reclassified; "rawstate" would give the original SPINDLE scoring
EPOCH_SEC = 4  # SPINDLE epoch length in seconds

# =====================================================================================
# END OF SETTINGS
# =====================================================================================

STATES = ["w", "n", "r"]
STATE_NAMES = {"w": "Wake", "n": "NREM", "r": "REM"}
METRICS = ["frac_time", "n_bouts", "mean_bout_sec"]
WHOLE_METRICS = METRICS + ["bouts_per_hour"]
SHEET_NAME = {"n_bouts": "bouts"}  # sheet names: w_frac_time, w_bouts, w_mean_bout_sec, ...


def stop(*problems):
    """Print what is wrong with the settings and quit."""
    print("\n" + "=" * 78)
    print("CANNOT RUN YET. Fix the following in the SETTINGS block of this script:\n")
    for problem in problems:
        print(f"  * {problem}")
    print("=" * 78)
    sys.exit(1)


def animal_id(path):
    return path.name.split("_")[0]  # the filename text before the first "_"


def edf_start(path):
    """Recording start from the EDF header (date dd.mm.yy at bytes 168-175, time hh.mm.ss at 176-183)."""
    with open(path, "rb") as f:
        header = f.read(256)
    return datetime.strptime(f"{header[168:176].decode()} {header[176:184].decode()}", "%d.%m.%y %H.%M.%S")


def read_groups_csv(path):
    """``{group: [animal IDs]}`` from a CSV, in either of two layouts.

    * One row per animal, with columns named ``animal`` and ``group`` (other columns are ignored).
    * One column per group: the group name in the first row, its animal IDs below it.

    Blank cells and rows are skipped, and a group with no animals is dropped. An animal listed
    in two groups is an error.
    """
    with open(path, newline="", encoding="utf-8-sig") as f:
        lines = [(n, [cell.strip() for cell in row]) for n, row in enumerate(csv.reader(f), start=1)]
    lines = [(n, row) for n, row in lines if any(row)]
    if not lines:
        raise ValueError(f"{path} is empty")
    header, body = lines[0][1], lines[1:]
    lowered = [name.lower() for name in header]

    groups = {}
    if "animal" in lowered and "group" in lowered:
        i_animal, i_group = lowered.index("animal"), lowered.index("group")
        for n, row in body:
            animal = row[i_animal] if i_animal < len(row) else ""
            group = row[i_group] if i_group < len(row) else ""
            if animal and group:
                groups.setdefault(group, []).append(animal)
            elif animal or group:
                raise ValueError(f"{path}, line {n}: a row needs both an animal and a group")
    else:
        for col, name in enumerate(header):
            animals = [row[col] for _, row in body if col < len(row) and row[col]]
            if animals and not name:
                raise ValueError(f"{path}: column {col + 1} has animals but no group name in its first row")
            if name:
                groups[name] = animals

    groups = {group: list(dict.fromkeys(animals)) for group, animals in groups.items() if animals}
    owner = {}
    for group, animals in groups.items():
        for animal in animals:
            if owner.setdefault(animal, group) != group:
                raise ValueError(f"{path}: {animal} is in more than one group ({owner[animal]}, {group})")
    return groups


def check_settings():
    """Stop with a plain message if a setting is still a placeholder or points at nothing."""
    problems = []
    if str(PROJECT_ROOT) == "/path/to/your/project":
        problems.append("PROJECT_ROOT is still the placeholder. Set it to your project folder.")
    elif not (Path(PROJECT_ROOT) / "manifest.json").is_file():
        problems.append(
            f"PROJECT_ROOT = {PROJECT_ROOT} has no manifest.json. "
            "Set it to the folder that holds manifest.json."
        )
    if GROUPS and GROUPS_CSV:
        problems.append("Both GROUPS and GROUPS_CSV are set. Use one of them (GROUPS_CSV wins otherwise).")
    if GROUPS_CSV and not Path(GROUPS_CSV).is_file():
        problems.append(f"GROUPS_CSV = {GROUPS_CSV} does not exist.")
    if TZ:
        try:
            ZoneInfo(TZ)
        except Exception:
            problems.append(f'TZ = "{TZ}" is not a time zone name (examples: "America/Chicago", "UTC").')
    if 24 % BIN_HOURS:
        problems.append(f"BIN_HOURS = {BIN_HOURS} must divide 24 (1, 2, 3, 4, 6, 8 or 12).")
    if not 0 <= BEGIN_HR < 24:
        problems.append(f"BEGIN_HR = {BEGIN_HR} must be an hour from 0 to 23.")
    if problems:
        stop(*problems)


def load_project():
    """Read the project: animals, start times, exclusions, groups. Stops on anything missing."""
    root = Path(PROJECT_ROOT)
    manifest = json.loads((root / "manifest.json").read_text())
    spindle_dir = root / "spindle_outputs" / manifest["combo"]

    csv_paths = sorted(spindle_dir.glob("*.csv"))
    if not csv_paths:
        stop(f"No state CSVs in {spindle_dir}. Has this project been through SPINDLE and Reclassify?")
    ids = [animal_id(p) for p in csv_paths]
    if len(ids) != len(set(ids)):
        stop(f"Two CSVs for one animal in {spindle_dir}.")

    edfs = {animal_id(p): p for p in (root / "standard").glob("*PROCESSED.edf")}
    missing = sorted(set(ids) - set(edfs))
    if missing:
        stop(f"No EDF in {root / 'standard'} for: {missing}")
    starts = {a: edf_start(edfs[a]) for a in ids}

    no_column = [p.name for p in csv_paths if STATE_COLUMN not in pd.read_csv(p, nrows=0).columns]
    if no_column:
        stop(
            f"These CSVs have no '{STATE_COLUMN}' column: {no_column}. "
            "Run the Reclassify step in ebb-viewer first (or set STATE_COLUMN)."
        )

    # animals excluded in Sleep QC / PSD review; the reason is in the animal's notes
    excluded = {
        a: e["notes"].strip().replace("\n", "; ")
        for a, e in manifest["animals"].items()
        if e["sleep_qc"] == "excluded"
    }
    undecided = [a for a in ids if manifest["animals"].get(a, {}).get("sleep_qc") not in ("ok", "excluded")]
    if undecided:
        print("WARNING: Sleep QC not decided for", undecided)
    for a, why in excluded.items():
        print(f"excluded: {a}: {why}")
    if not INCLUDE_EXCLUDED:
        csv_paths = [p for p in csv_paths if animal_id(p) not in excluded]
        ids = [animal_id(p) for p in csv_paths]

    groups = read_groups_csv(GROUPS_CSV) if GROUPS_CSV else (GROUPS or {"All": sorted(ids)})
    group_of = {a: g for g, members in groups.items() for a in members}
    if len(group_of) != sum(map(len, groups.values())):
        stop("An animal is in two groups. Check GROUPS / GROUPS_CSV.")
    ungrouped = sorted(set(ids) - set(group_of))
    if ungrouped:
        stop(f"These animals are in the project but in no group: {ungrouped}. Add them to GROUPS / GROUPS_CSV.")
    unknown = sorted(set(group_of) - set(ids))
    if unknown:
        print(f"WARNING: in your groups but not analysed (excluded, or no data): {unknown}")
    print(f"\n{len(ids)} animals analysed:", {g: len([a for a in m if a in ids]) for g, m in groups.items()})
    return csv_paths, starts, excluded, groups, group_of


# ------------------------------------- calculations -------------------------------------


def load_states(path):
    """One letter (w / n / r) per 4 s epoch."""
    states = pd.read_csv(path, usecols=[STATE_COLUMN])[STATE_COLUMN].to_numpy()
    assert set(states) <= set(STATES), f"unexpected states in {path.name}: {set(states) - set(STATES)}"
    return states


def bouts(mask):
    """(number of bouts, mean bout length in samples) for a boolean mask; a bout is an unbroken run."""
    change = np.diff(mask.astype(int))
    starts_, ends_ = np.where(change == 1)[0] + 1, np.where(change == -1)[0] + 1
    if mask[0]:
        starts_ = np.r_[0, starts_]
    if mask[-1]:
        ends_ = np.r_[ends_, len(mask)]
    return (len(starts_), (ends_ - starts_).mean()) if len(starts_) else (0, np.nan)


def seconds_until_next(start, hour, tz):
    """Real seconds from `start` to the next `hour`:00 (compared in UTC so daylight saving counts)."""
    if tz:
        start = start.replace(tzinfo=tz)
    target = start.replace(hour=hour, minute=0, second=0, microsecond=0)
    if start > target:
        target = datetime.combine((start + timedelta(days=1)).date(), time(hour), tzinfo=start.tzinfo)
    if tz:
        start, target = start.astimezone(timezone.utc), target.astimezone(timezone.utc)
    return int((target - start).total_seconds())


def sleep_bins(path, starts, tz):
    """Per animal: rotate to BEGIN_HR, cut into BIN_HOURS bins, score every state in every bin."""
    animal = animal_id(path)
    secs = np.repeat(load_states(path), EPOCH_SEC)  # one state per second
    offset = seconds_until_next(starts[animal], BEGIN_HR, tz)
    if offset >= len(secs):
        print(f"skipping {animal}: recording ends before it reaches {BEGIN_HR}:00")
        return None
    rotated = np.concatenate([secs[offset:], secs[:offset]])  # first BEGIN_HR is position 0
    rows = []
    for i, s in enumerate(range(0, len(rotated), BIN_HOURS * 3600)):
        chunk = rotated[s : s + BIN_HOURS * 3600]
        for state in STATES:
            m = chunk == state
            n, mean_len = bouts(m)
            rows.append(
                {
                    "mouse_id": animal,
                    "bin_idx": i,
                    "bin_start_hr": BEGIN_HR + i * BIN_HOURS,
                    "state": state,
                    "frac_time": m.sum() / len(chunk),
                    "n_bouts": n,
                    "mean_bout_sec": mean_len,
                }
            )
    return pd.DataFrame(rows)


def whole_file(path):
    """Per animal and state over the whole recording (no rotation)."""
    states = load_states(path)
    hours = len(states) * EPOCH_SEC / 3600
    rows = []
    for state in STATES:
        n, mean_epochs = bouts(states == state)
        rows.append(
            {
                "mouse_id": animal_id(path),
                "state": state,
                "frac_time": (states == state).mean(),
                "n_bouts": n,
                "mean_bout_sec": mean_epochs * EPOCH_SEC,
                "bouts_per_hour": n / hours,
                "n_hours": hours,
                "n_days": hours / 24,
            }
        )
    return pd.DataFrame(rows)


def calculate(csv_paths, starts, tz, excluded, group_of):
    bins = pd.concat([d for d in (sleep_bins(p, starts, tz) for p in csv_paths) if d is not None], ignore_index=True)

    # fold onto a 24 h day (hour 30 = 6 AM again), average across days, convert to ZT
    bins["hour_of_day"] = (bins["bin_start_hr"] % 24).astype(int)
    g = bins.groupby(["mouse_id", "state", "hour_of_day"], as_index=False)
    by_hour = g[METRICS].mean()
    by_hour["n_days"] = g.size()["size"]
    by_hour["zt"] = (by_hour["hour_of_day"] - BEGIN_HR) % 24
    whole = pd.concat([whole_file(p) for p in csv_paths], ignore_index=True)

    for table in (by_hour, whole):
        table["group"] = table["mouse_id"].map(group_of)
        table["excluded"] = table["mouse_id"].isin(excluded)
    by_hour = by_hour.sort_values(["mouse_id", "state", "zt"], ignore_index=True)
    print(f"{bins['mouse_id'].nunique()} animals, {len(bins)} bin rows, {len(by_hour)} by-hour rows")
    return bins, by_hour, whole


# ------------------------------- tables and Prism workbooks -------------------------------


def animal_columns(df, groups):
    """Animals blocked by group, each block padded with blanks to the size of the largest group."""
    present = set(df["mouse_id"])
    blocks = {gname: sorted(a for a in members if a in present) for gname, members in groups.items()}
    width = max(map(len, blocks.values()))
    return [
        (gname, a) for gname, block in blocks.items() if block for a in block + [None] * (width - len(block))
    ]


def write_workbook(path, sheets, layout):
    """`sheets`: {sheet name: (row labels, wide table with one column per animal, top-left label)}."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, (labels, wide, corner) in sheets.items():
        ws = wb.create_sheet(name)
        ws.cell(row=2, column=1, value=corner)
        for c, (gname, animal) in enumerate(layout, start=2):
            ws.cell(row=1, column=c, value=gname)
            ws.cell(row=2, column=c, value=animal)
        for r, label in enumerate(labels, start=3):
            ws.cell(row=r, column=1, value=label)
            for c, (_, animal) in enumerate(layout, start=2):
                val = wide.at[wide.index[r - 3], animal] if animal in wide.columns else None
                ws.cell(row=r, column=c, value=None if pd.isna(val) else float(val))
        ws.freeze_panes = "B3"
    wb.save(path)


def write_tables(out_dir, bins, by_hour, whole, groups, zt_label):
    layout = animal_columns(by_hour, groups)

    by_hour_sheets = {}
    for metric in METRICS:
        for state in STATES:
            wide = by_hour[by_hour["state"] == state].pivot(index="zt", columns="mouse_id", values=metric).sort_index()
            labels = [f"{zt}-{zt + BIN_HOURS}" for zt in wide.index]
            by_hour_sheets[f"{state}_{SHEET_NAME.get(metric, metric)}"] = (labels, wide, zt_label)
    write_workbook(out_dir / "prism_sleep_by_hour.xlsx", by_hour_sheets, layout)

    whole_sheets = {}
    for metric in WHOLE_METRICS:
        wide = whole.pivot(index="state", columns="mouse_id", values=metric).reindex(STATES)
        whole_sheets[SHEET_NAME.get(metric, metric)] = ([STATE_NAMES[s] for s in STATES], wide, "State")
    write_workbook(out_dir / "prism_sleep_whole_file.xlsx", whole_sheets, layout)

    bins.to_csv(out_dir / "sleep_bins.csv", index=False)
    by_hour.to_csv(out_dir / "sleep_bins_avg.csv", index=False)
    whole.to_csv(out_dir / "sleep_whole_file.csv", index=False)


# ------------------------------------------ plots ------------------------------------------


def plot_by_hour(out_dir, by_hour, groups, zt_label):
    """Group mean +/- SEM per ZT bin (grey band = lights-off)."""
    rows = [
        ("n_bouts", "Bout number", "bouts per bin", 1),
        ("mean_bout_sec", "Bout duration", "seconds", 1),
        ("frac_time", "Time in state", "% of bin", 100),
    ]
    fig, axes = plt.subplots(3, 3, figsize=(12, 9), sharex=True)
    for r, (metric, title, ylabel, scale) in enumerate(rows):
        for c, state in enumerate(STATES):
            ax = axes[r, c]
            ax.axvspan(DARK_ONSET_ZT, 24, color="0.88", zorder=0)
            for gname in groups:
                d = by_hour[(by_hour["state"] == state) & (by_hour["group"] == gname)]
                agg = d.groupby("zt")[metric].agg(["mean", "sem"])
                ax.errorbar(agg.index, agg["mean"] * scale, yerr=agg["sem"] * scale, marker="o", capsize=3, label=gname)
            ax.set_title(f"{STATE_NAMES[state]}: {title}")
            ax.set_xlim(0, 24)
            ax.set_xticks(range(0, 25, 4))
            if c == 0:
                ax.set_ylabel(ylabel)
            if r == 2:
                ax.set_xlabel(zt_label)
    axes[0, 0].legend()
    fig.tight_layout()
    fig.savefig(out_dir / "sleep_by_hour.png", dpi=150)
    plt.close(fig)


def plot_whole_file(out_dir, whole, groups):
    """Boxes with every animal as a point."""
    titles = {
        "frac_time": ("Time in state", "% of recording", 100),
        "n_bouts": ("Bout number", "bouts per recording", 1),
        "mean_bout_sec": ("Bout duration", "seconds", 1),
        "bouts_per_hour": ("Bout rate", "bouts per hour", 1),
    }
    rng = np.random.default_rng(0)  # fixed seed: same jitter every run
    names = list(groups)
    box_w = 0.8 / len(names)
    fig, axes = plt.subplots(1, 4, figsize=(16, 3.8))
    for ax, metric in zip(axes, WHOLE_METRICS):
        title, ylabel, scale = titles[metric]
        for gi, gname in enumerate(names):
            for si, state in enumerate(STATES):
                y = whole.loc[(whole["state"] == state) & (whole["group"] == gname), metric].dropna() * scale
                x = si + (gi - (len(names) - 1) / 2) * box_w
                ax.boxplot(
                    y,
                    positions=[x],
                    widths=box_w * 0.8,
                    showfliers=False,
                    medianprops={"color": f"C{gi}"},
                    boxprops={"color": f"C{gi}"},
                    whiskerprops={"color": f"C{gi}"},
                    capprops={"color": f"C{gi}"},
                )
                ax.scatter(
                    x + rng.uniform(-box_w / 4, box_w / 4, len(y)),
                    y,
                    s=14,
                    color=f"C{gi}",
                    zorder=3,
                    label=gname if si == 0 else None,
                )
        ax.set_xticks(range(len(STATES)), [STATE_NAMES[s] for s in STATES])
        ax.set_title(title)
        ax.set_ylabel(ylabel)
    axes[0].legend()
    fig.tight_layout()
    fig.savefig(out_dir / "sleep_whole_file.png", dpi=150)
    plt.close(fig)


def main():
    check_settings()
    out_dir = Path(OUT_DIR) if OUT_DIR else Path(PROJECT_ROOT) / "sleep_analysis"
    print("=" * 78)
    print("RESULTS WILL BE SAVED IN:\n    ", out_dir.resolve())
    print("=" * 78)

    tz = ZoneInfo(TZ) if TZ else None
    csv_paths, starts, excluded, groups, group_of = load_project()
    out_dir.mkdir(parents=True, exist_ok=True)

    bins, by_hour, whole = calculate(csv_paths, starts, tz, excluded, group_of)
    zt_label = f"ZT hours (0 = {BEGIN_HR % 12 or 12} {'AM' if BEGIN_HR % 24 < 12 else 'PM'})"
    write_tables(out_dir, bins, by_hour, whole, groups, zt_label)
    plot_by_hour(out_dir, by_hour, groups, zt_label)
    plot_whole_file(out_dir, whole, groups)

    print("\nWrote:", *sorted(p.name for p in out_dir.iterdir()), sep="\n  ")
    print("=" * 78)
    print("ALL RESULTS ARE IN:\n    ", out_dir.resolve())
    print("=" * 78)


if __name__ == "__main__":
    main()
