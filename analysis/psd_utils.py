"""Helpers for ``psd_analysis.ipynb`` and ``psd_analysis.py``: turn a sealed ``psds_final.pkl`` into
tables and plots.

Keep this file in the same folder as the notebook or script. Edit the notebook's Settings cell (or
the script's SETTINGS block) and run it. The functions are small and each does one step, so the notebook reads as a
recipe: load -> assign groups -> interpolate the line noise -> normalize -> plot -> band power ->
export for Prism.

Conventions (these match the ebb pipeline's output):

* ``psds_final.pkl`` is a list of dicts, one per animal and state, each with a ``(channels, freqs)``
  array ``psd`` and the matching ``freqs``. Excluded animals carry ``excluded = True``.
* A long-format DataFrame has one row per animal x state x channel, with the spectrum in a column
  of 1-D arrays. The frequency axis is stored once, in ``df.attrs["freqs"]``.
* Power is in uV^2/Hz. Normalized spectra are in arbitrary units (a.u.).
"""

from __future__ import annotations

import csv
import pathlib
import pickle
import warnings
from collections.abc import Mapping, Sequence
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from specparam.utils import interpolate_spectrum

# Seaborn's boxplot still passes the deprecated ``vert=`` kwarg to matplotlib. Harmless, but noisy
# in loops, so silence just that warning.
warnings.filterwarnings(
    "ignore",
    message=".*vert: bool was deprecated.*",
    category=mpl.MatplotlibDeprecationWarning,
)

CHANNEL_NAMES = {0: "LFC", 1: "RSC", 2: "LSC"}
STATES = ("wake", "nrem", "rem")

# The recording hardware attenuates the lowest frequencies. These factors (for 0.5, 1.0, ..., 3.5 Hz)
# correct the PSD there. ``records_to_df`` applies them unless told not to.
INVERSE_FILTER_RANGE = (0.5, 3.5)
INVERSE_FILTER = np.array(
    [2.485297985, 1.360646853, 1.145286775, 1.071699399, 1.038017048, 1.018656356, 1.007806318]
)

DEFAULT_BANDS = {
    "1-4 Hz": (1, 4),
    "4-8 Hz": (4, 8),
    "8-12 Hz": (8, 12),
    "12-30 Hz": (12, 30),
    "30-70 Hz": (30, 70),
    "70-100 Hz": (70, 100),
}

# Line noise bridged by interpolation, in Hz. The pipeline applies a 60 Hz notch filter by default
# (stopband 57-63 Hz), so 56-64 Hz is bridged, as in the PSD Visualizer. If the run was computed
# without the notch, 58-62 Hz is enough. The notch does not touch 120 Hz.
DEFAULT_INTERP_RANGES = ((56, 64), (118, 122))

_ID_COLUMNS = ("animal", "time", "state", "group", "ch", "channel")


# ----------------------------------------------------------------------------------------------
# Loading
# ----------------------------------------------------------------------------------------------
def load_psd_records(path: str | Path) -> list[dict]:
    """Read a ``psds_final.pkl`` (or any PSD pickle) and return its list of records.

    Files written on Windows hold ``WindowsPath`` objects, which cannot be unpickled on macOS or
    Linux; this reads them as ordinary paths instead.
    """
    path = Path(path)
    try:
        with open(path, "rb") as f:
            records = pickle.load(f)
    except NotImplementedError:
        pathlib.WindowsPath = pathlib.PosixPath  # type: ignore[misc]
        with open(path, "rb") as f:
            records = pickle.load(f)
    if isinstance(records, dict):
        records = [records]
    return list(records)


def default_output_dir(*psds_final: str | Path) -> Path:
    """Where results go when ``OUT_DIR`` is not set: a folder per analysed run.

    For ``<project>/psd/<run>/psds_final.pkl`` that is ``<project>/psd_analysis/<run>``, so two
    runs of the same project do not overwrite each other. For a pickle anywhere else it is
    ``psd_analysis/<file name without .pkl>`` beside the file. The project (or folder) is the
    first file's. Given several files (several time points) the folder is named after all of
    them, joined with ``+``. The notebook's own folder is never used.
    """
    paths = [Path(p).expanduser().resolve() for p in psds_final]
    first = paths[0]
    in_project = first.parent.parent.name == "psd"
    base = (first.parent.parent.parent if in_project else first.parent) / "psd_analysis"
    names = [p.parent.name if p.parent.parent.name == "psd" else p.stem for p in paths]
    return base / "+".join(dict.fromkeys(names))


def get_freqs(df: pd.DataFrame) -> np.ndarray:
    """The frequency axis (Hz) shared by every spectrum in ``df``."""
    freqs = df.attrs.get("freqs")
    if freqs is None:
        raise ValueError(
            "This DataFrame has no frequency axis (df.attrs['freqs']). Build it with "
            "records_to_df or load_psd_runs, and avoid operations that drop df.attrs."
        )
    return np.asarray(freqs, dtype=float)


def records_to_df(
    records: Sequence[Mapping],
    *,
    apply_inverse_filter: bool = True,
    include_excluded: bool = False,
    channel_names: Mapping[int, str] = CHANNEL_NAMES,
) -> pd.DataFrame:
    """Long-format table (one row per animal x state x channel) from PSD records.

    Columns: ``animal``, ``state``, ``ch``, ``channel``, ``spectra`` (1-D array, uV^2/Hz),
    ``excluded``, ``exclusion_reason`` (which includes the step: Sleep QC or PSD review),
    ``n_estimates`` and the fraction of samples masked by each source. The animal ID
    is the part of the EDF file name before the first ``_``.

    Args:
        records: the list from :func:`load_psd_records`.
        apply_inverse_filter: multiply 0.5-3.5 Hz by :data:`INVERSE_FILTER` (needs 0.5 Hz bins).
            The records themselves are not modified.
        include_excluded: keep animals excluded in Sleep QC or PSD review (marked in the
            ``excluded`` column). By default they are left out and listed with their reasons.
        channel_names: label for each channel index.
    """
    if not records:
        raise ValueError("No PSD records to convert.")
    freqs = np.asarray(records[0]["freqs"], dtype=float)
    inv_mask = (freqs >= INVERSE_FILTER_RANGE[0]) & (freqs <= INVERSE_FILTER_RANGE[1])
    if apply_inverse_filter and inv_mask.sum() != len(INVERSE_FILTER):
        raise ValueError(
            f"The inverse filter needs {len(INVERSE_FILTER)} frequency bins between "
            f"{INVERSE_FILTER_RANGE[0]} and {INVERSE_FILTER_RANGE[1]} Hz (0.5 Hz resolution), "
            f"found {int(inv_mask.sum())}. Set apply_inverse_filter=False."
        )

    rows = []
    for rec in records:
        if not np.array_equal(np.asarray(rec["freqs"], dtype=float), freqs):
            raise ValueError("The records do not share one frequency axis.")
        animal = Path(rec["path"]).stem.split("_")[0]
        for ch, psd in enumerate(np.asarray(rec["psd"], dtype=float)):
            spectrum = psd.copy()
            if apply_inverse_filter:
                spectrum[inv_mask] = spectrum[inv_mask] * INVERSE_FILTER
            rows.append(
                {
                    "animal": animal,
                    "state": rec["behavior"][-1],
                    "ch": ch,
                    "channel": channel_names.get(ch, f"ch{ch}"),
                    "spectra": spectrum,
                    "excluded": bool(rec.get("excluded", False)),
                    "exclusion_reason": rec.get("exclusion_reason", ""),
                    "n_estimates": rec.get("n_estimates"),
                    "spindle_out_frac": rec.get("spindle_out_frac"),
                    "comm_out_frac": rec.get("comm_out_frac"),
                    "sd_out_frac": rec.get("sd_out_frac"),
                    "swd_out_frac": rec.get("swd_out_frac"),
                    "ann_out_frac": rec.get("ann_out_frac"),
                }
            )
    df = pd.DataFrame(rows)
    if not include_excluded:
        dropped = df[df["excluded"]].drop_duplicates("animal").sort_values("animal")
        if len(dropped):
            print(
                f"Left out {len(dropped)} excluded animal(s) (reason, with the step that set it):"
            )
            for _, row in dropped.iterrows():
                print(f"  {row['animal']}: {row['exclusion_reason'] or 'no reason recorded'}")
        df = df[~df["excluded"]].reset_index(drop=True)
    df.attrs["freqs"] = tuple(freqs.tolist())  # a tuple: pandas compares attrs with ==
    return df


def load_psd_runs(
    runs: Mapping[str, str | Path],
    *,
    apply_inverse_filter: bool = True,
    include_excluded: bool = False,
    channel_names: Mapping[int, str] = CHANNEL_NAMES,
) -> pd.DataFrame:
    """Load one or more ``psds_final.pkl`` files into one table with a ``time`` column.

    ``runs`` maps a label to a file, for example ``{"pre": ".../psds_final.pkl"}``. Use one entry
    for a single time point, or several (``pre`` / ``post``) to compare them.
    """
    frames = []
    for label, path in runs.items():
        part = records_to_df(
            load_psd_records(path),
            apply_inverse_filter=apply_inverse_filter,
            include_excluded=include_excluded,
            channel_names=channel_names,
        )
        part.insert(0, "time", label)
        frames.append(part)
    freqs = get_freqs(frames[0])
    for part in frames[1:]:
        if not np.array_equal(get_freqs(part), freqs):
            raise ValueError("The runs do not share one frequency axis.")
        part.attrs = {}  # the axis is set again on the combined table
    frames[0].attrs = {}
    df = pd.concat(frames, ignore_index=True)
    df.attrs["freqs"] = tuple(freqs.tolist())
    return df


# ----------------------------------------------------------------------------------------------
# Groups, line noise, normalization, band power
# ----------------------------------------------------------------------------------------------
def read_groups_csv(path):
    """``{group: [animal IDs]}`` from a CSV, in either of two layouts.

    * One row per animal, with columns named ``animal`` and ``group`` (other columns are ignored).
    * One column per group: the group name in the first row, its animal IDs below it.

    Blank cells and rows are skipped, and a group with no animals is dropped. An animal listed
    in two groups is an error.
    """
    with open(path, newline="", encoding="utf-8-sig") as f:
        lines = [
            (n, [cell.strip() for cell in row]) for n, row in enumerate(csv.reader(f), start=1)
        ]
    lines = [(n, row) for n, row in lines if any(row)]
    if not lines:
        raise ValueError(f"{path} is empty")
    header, body = lines[0][1], lines[1:]
    lowered = [name.lower() for name in header]

    groups = {}
    if "animal" in lowered and "group" in lowered:
        i_animal, i_group = lowered.index("animal"), lowered.index("group")
        for n, row in body:
            animal = row[i_animal] if i_animal < len(row) else ""
            group = row[i_group] if i_group < len(row) else ""
            if animal and group:
                groups.setdefault(group, []).append(animal)
            elif animal or group:
                raise ValueError(f"{path}, line {n}: a row needs both an animal and a group")
    else:
        for col, name in enumerate(header):
            animals = [row[col] for _, row in body if col < len(row) and row[col]]
            if animals and not name:
                raise ValueError(
                    f"{path}: column {col + 1} has animals but no group name in its first row"
                )
            if name:
                groups[name] = animals

    groups = {group: list(dict.fromkeys(animals)) for group, animals in groups.items() if animals}
    owner = {}
    for group, animals in groups.items():
        for animal in animals:
            if owner.setdefault(animal, group) != group:
                raise ValueError(
                    f"{path}: {animal} is in more than one group ({owner[animal]}, {group})"
                )
    return groups


def assign_groups(df: pd.DataFrame, groups: Mapping[str, Sequence[str]]) -> pd.DataFrame:
    """Add ``group`` (and ``group_time`` when there is a ``time`` column).

    ``groups`` maps a group name to its animal IDs. Animals in no group are labelled
    ``unassigned``; listed animals that are not in the data are reported.
    """
    lookup = {animal: group for group, animals in groups.items() for animal in animals}
    out = df.copy()
    out.attrs = dict(df.attrs)
    out["group"] = out["animal"].map(lookup).fillna("unassigned")
    if "time" in out.columns:
        out["group_time"] = out["group"] + "_" + out["time"].astype(str)

    unassigned = sorted(out.loc[out["group"] == "unassigned", "animal"].unique())
    if unassigned:
        print(f"Not in any group ({len(unassigned)}): {', '.join(unassigned)}")
    missing = sorted(set(lookup) - set(out["animal"]))
    if missing:
        print(f"In GROUPS but not in the data ({len(missing)}): {', '.join(missing)}")
    return out


def interpolate_line_noise(
    df: pd.DataFrame,
    ranges: Sequence[Sequence[float]] = DEFAULT_INTERP_RANGES,
    source: str = "spectra",
    out: str = "interp_spectra",
) -> pd.DataFrame:
    """Add ``out``: ``source`` with each frequency range in ``ranges`` bridged by interpolation."""
    freqs = get_freqs(df)
    result = df.copy()
    result.attrs = dict(df.attrs)
    result[out] = df[source].apply(
        lambda psd: interpolate_spectrum(freqs, psd, [list(r) for r in ranges])[1]
    )
    return result


def compute_bandpower(spectrum: np.ndarray, freqs: np.ndarray, band: Sequence[float]) -> float:
    """Power in ``band`` (Hz): the area under the spectrum, by the trapezoid rule."""
    mask = (freqs >= band[0]) & (freqs <= band[1])
    return float(np.trapezoid(spectrum[mask], freqs[mask]))


def normalize_spectra(
    df: pd.DataFrame,
    source: str = "interp_spectra",
    out: str = "norm_spectra",
    band: Sequence[float] = (1, 100),
) -> pd.DataFrame:
    """Add ``out``: each spectrum divided by its own power in ``band`` (so that power is 1)."""
    freqs = get_freqs(df)
    result = df.copy()
    result.attrs = dict(df.attrs)
    result[out] = df[source].apply(lambda psd: psd / compute_bandpower(psd, freqs, band))
    return result


def band_power_table(
    df: pd.DataFrame,
    column: str,
    bands: Mapping[str, Sequence[float]] = DEFAULT_BANDS,
) -> pd.DataFrame:
    """One row per animal x state x channel (x time) with a column of power per band."""
    freqs = get_freqs(df)
    keep = [c for c in _ID_COLUMNS if c in df.columns]
    table = df[keep].copy()
    for name, band in bands.items():
        table[name] = [compute_bandpower(psd, freqs, band) for psd in df[column]]
    return table.reset_index(drop=True)


# ----------------------------------------------------------------------------------------------
# Plots
# ----------------------------------------------------------------------------------------------
def _present(df: pd.DataFrame, states: Sequence[str]) -> list[str]:
    return [s for s in states if s in set(df["state"])]


def _ylabel(column: str, log_y: bool) -> str:
    label = (
        "Normalized power [a.u.]" if "norm" in column.lower() else r"Power [$\mu V^2/\mathrm{Hz}$]"
    )
    return ("Log " + label) if log_y else label


def _finish(fig: plt.Figure, save_path: str | Path | None) -> plt.Figure:
    fig.tight_layout()
    if save_path is not None:
        Path(save_path).parent.mkdir(parents=True, exist_ok=True)
        fig.savefig(save_path, dpi=300, bbox_inches="tight")
    return fig


def plot_spectra(
    df: pd.DataFrame,
    column: str = "interp_spectra",
    *,
    group_col: str = "group",
    color_map: Mapping[str, str] | None = None,
    freq_range: Sequence[float] = (0.5, 110),
    log_y: bool = True,
    show_individuals: bool = False,
    use_median: bool = False,
    sharey: bool = True,
    title: str = "Spectra",
    states: Sequence[str] = STATES,
    channel_names: Mapping[int, str] = CHANNEL_NAMES,
    save_path: str | Path | None = None,
) -> plt.Figure:
    """Group spectra for each channel (rows) and state (columns).

    Draws each group's mean (or median) with a standard-error band; ``show_individuals`` draws
    every animal as a faint line instead. ``color_map`` maps group -> color and sets the order.
    """
    freqs_all = get_freqs(df)
    mask = (freqs_all >= freq_range[0]) & (freqs_all <= freq_range[1])
    freqs = freqs_all[mask]
    states = _present(df, states)
    channels = sorted(df["ch"].unique())
    if color_map is None:
        color_map = {g: plt.cm.tab10(i) for i, g in enumerate(df[group_col].unique())}

    fig, axs = plt.subplots(
        len(channels),
        len(states),
        figsize=(3.4 * len(states), 2.7 * len(channels)),
        sharey=sharey,
        squeeze=False,
    )
    for ci, ch in enumerate(channels):
        for si, state in enumerate(states):
            ax = axs[ci][si]
            cell = df[(df["state"] == state) & (df["ch"] == ch)]
            for group, color in color_map.items():
                psds = [p[mask] for p in cell.loc[cell[group_col] == group, column]]
                if not psds:
                    continue
                psds = np.stack(psds)
                n = len(psds)
                if show_individuals:
                    for i, psd in enumerate(psds):
                        ax.plot(
                            freqs, psd, color=color, alpha=0.25, label=group if i == 0 else None
                        )
                    continue
                center = np.median(psds, axis=0) if use_median else np.mean(psds, axis=0)
                ax.plot(
                    freqs, center, color=color, alpha=0.8, linewidth=1.3, label=f"{group} ({n})"
                )
                if n > 1:
                    se = np.std(psds, axis=0, ddof=1) / np.sqrt(n)
                    ax.fill_between(freqs, center - se, center + se, alpha=0.12, color=color)
            if log_y:
                ax.set_yscale("log")
            ax.grid(alpha=0.2)
            ax.spines[["top", "right"]].set_visible(False)
            if ci == 0:
                ax.set_title(state.upper())
            if si == 0:
                ax.set_ylabel(channel_names.get(ch, f"ch{ch}"))
            if ci == 0 and si == 0:
                ax.legend(fontsize="x-small", loc="upper right", frameon=False)
    fig.supylabel(_ylabel(column, log_y))
    fig.supxlabel("Frequency [Hz]")
    fig.suptitle(title)
    return _finish(fig, save_path)


def plot_spectra_by_animal(
    df: pd.DataFrame,
    column: str = "interp_spectra",
    *,
    freq_range: Sequence[float] = (0.5, 110),
    log_y: bool = True,
    sharey: bool = True,
    title: str = "Spectra by animal",
    states: Sequence[str] = STATES,
    channel_names: Mapping[int, str] = CHANNEL_NAMES,
    save_path: str | Path | None = None,
) -> plt.Figure:
    """One line per animal, for each channel (rows) and state (columns)."""
    freqs_all = get_freqs(df)
    mask = (freqs_all >= freq_range[0]) & (freqs_all <= freq_range[1])
    freqs = freqs_all[mask]
    states = _present(df, states)
    channels = sorted(df["ch"].unique())
    animals = list(df["animal"].unique())
    colors = {a: plt.cm.tab20(i % 20) for i, a in enumerate(animals)}

    fig, axs = plt.subplots(
        len(channels),
        len(states),
        figsize=(3.4 * len(states), 2.7 * len(channels)),
        sharey=sharey,
        squeeze=False,
    )
    for ci, ch in enumerate(channels):
        for si, state in enumerate(states):
            ax = axs[ci][si]
            cell = df[(df["state"] == state) & (df["ch"] == ch)]
            for _, row in cell.iterrows():
                style = "-" if animals.index(row["animal"]) < 20 else "--"
                ax.plot(
                    freqs,
                    row[column][mask],
                    color=colors[row["animal"]],
                    linestyle=style,
                    linewidth=1,
                    alpha=0.8,
                    label=row["animal"],
                )
            if log_y:
                ax.set_yscale("log")
            ax.grid(alpha=0.2)
            ax.spines[["top", "right"]].set_visible(False)
            if ci == 0:
                ax.set_title(state.upper())
            if si == 0:
                ax.set_ylabel(channel_names.get(ch, f"ch{ch}"))
    handles, labels = axs[0][0].get_legend_handles_labels()
    seen = dict(zip(labels, handles, strict=True))
    fig.legend(
        seen.values(), seen.keys(), loc="upper left", bbox_to_anchor=(1, 0.95), fontsize="x-small"
    )
    fig.supylabel(_ylabel(column, log_y))
    fig.supxlabel("Frequency [Hz]")
    fig.suptitle(title)
    return _finish(fig, save_path)


def _outlier_mask(values: pd.Series, thresh: float = 1.5) -> pd.Series:
    """Tukey rule: True where a value is outside [Q1 - thresh*IQR, Q3 + thresh*IQR]."""
    x = values.astype(float)
    q1, q3 = np.nanpercentile(x, 25), np.nanpercentile(x, 75)
    iqr = q3 - q1
    return (x < q1 - thresh * iqr) | (x > q3 + thresh * iqr)


def boxplot_band_power(
    table: pd.DataFrame,
    band: str,
    ylabel: str,
    title: str,
    *,
    x: str = "group",
    hue: str | None = None,
    order: Sequence[str] | None = None,
    hue_order: Sequence[str] | None = None,
    iqr_thresh: float = 1.5,
    ylim: Sequence[float] | None = None,
    outlier_groups: Sequence[str] = ("state", "ch", "group"),
    states: Sequence[str] = STATES,
    channel_names: Mapping[int, str] = CHANNEL_NAMES,
    annotate_fontsize: float = 7,
    save_path: str | Path | None = None,
) -> tuple[plt.Figure, np.ndarray]:
    """Box plots with every animal as a point, for each state (rows) and channel (columns).

    Points outside the Tukey fences (computed within each of ``outlier_groups``) are outlined and
    labelled with the animal ID. ``table`` comes from :func:`band_power_table`; ``band`` names the
    column to plot. ``hue`` defaults to ``x``.
    """
    hue = hue or x
    df = table.copy()
    missing = {"state", "ch", "animal", x, hue, band} - set(df.columns)
    if missing:
        raise ValueError(f"Missing required columns: {sorted(missing)}")

    df["_is_outlier"] = False
    for _, part in df.groupby(list(outlier_groups)):
        df.loc[part.index, "_is_outlier"] = _outlier_mask(part[band], iqr_thresh)

    order = list(order) if order is not None else list(pd.unique(df[x]))
    hue_order = list(hue_order) if hue_order is not None else list(pd.unique(df[hue]))
    states = _present(df, states)
    channels = sorted(df["ch"].unique())

    fig, axs = plt.subplots(
        len(states),
        len(channels),
        figsize=(4 * len(channels), 2.7 * len(states)),
        sharex=True,
        squeeze=False,
    )
    for si, state in enumerate(states):
        for ci, ch in enumerate(channels):
            ax = axs[si][ci]
            cell = df[(df["state"] == state) & (df["ch"] == ch)]
            sns.boxplot(
                data=cell,
                x=x,
                y=band,
                hue=hue,
                order=order,
                hue_order=hue_order,
                fill=False,
                gap=0.05,
                linewidth=0.5,
                showfliers=False,
                ax=ax,
                legend=False,
            )
            sns.stripplot(
                data=cell,
                x=x,
                y=band,
                hue=hue,
                order=order,
                hue_order=hue_order,
                dodge=True,
                jitter=False,
                alpha=0.65,
                size=4.5,
                ax=ax,
                legend=True,
            )

            outliers = cell[cell["_is_outlier"]]
            if not outliers.empty:
                n_before = len(ax.collections)
                sns.stripplot(
                    data=outliers,
                    x=x,
                    y=band,
                    hue=hue,
                    order=order,
                    hue_order=hue_order,
                    dodge=True,
                    jitter=False,
                    size=5.5,
                    linewidth=0.8,
                    edgecolor="black",
                    ax=ax,
                    legend=False,
                )
                # Seaborn draws categories in order, then hues; label the points in that order.
                blocks = [
                    outliers[(outliers[x] == xc) & (outliers[hue] == hc)]
                    for xc in order
                    for hc in hue_order
                ]
                blocks = [b for b in blocks if not b.empty]
                ordered = pd.concat(blocks) if blocks else outliers
                offsets = [
                    off
                    for coll in ax.collections[n_before:]
                    if isinstance(coll, mpl.collections.PathCollection)
                    for off in coll.get_offsets()
                ]
                for (px, py), animal in zip(offsets, ordered["animal"], strict=False):
                    ax.annotate(
                        str(animal),
                        xy=(px, py),
                        xytext=(2, 0),
                        textcoords="offset points",
                        fontsize=annotate_fontsize,
                        rotation=25,
                        ha="left",
                        va="center",
                        bbox=dict(boxstyle="round,pad=0.15", fc="white", ec="none", alpha=0.7),
                    )

            ax.set_ylabel(f"{ylabel} {state}" if ci == 0 else "")
            ax.set_xlabel("")
            ax.set_title(f"{state} {channel_names.get(ch, f'ch{ch}')}")
            if ylim is not None:
                ax.set_ylim(ylim)
            if ax.get_legend():
                ax.get_legend().remove()
            ax.tick_params(axis="x", rotation=90)

    handles, labels = axs[-1][-1].get_legend_handles_labels()
    unique = {lab: h for h, lab in zip(handles, labels, strict=True) if lab != "_nolegend_"}
    if unique:
        fig.legend(
            unique.values(), unique.keys(), loc="upper left", bbox_to_anchor=(1, 1), title=hue
        )
    fig.suptitle(title)
    _finish(fig, save_path)
    return fig, axs


# ----------------------------------------------------------------------------------------------
# Export for GraphPad Prism
# ----------------------------------------------------------------------------------------------
def _sort_animals(rows: pd.DataFrame, sex_map: Mapping[str, str] | None) -> pd.DataFrame:
    """Order a group's animals alphabetically, or males first then females when ``sex_map`` is given."""
    if sex_map is None:
        return rows.sort_values("animal")
    rank = {"M": 0, "F": 1}
    key = rows["animal"].map(lambda a: (rank.get(str(sex_map.get(a, "")).upper(), 2), a))
    return rows.assign(_key=key).sort_values("_key").drop(columns="_key")


def _write_prism(
    df: pd.DataFrame,
    groups: Sequence[str],
    output_path: str | Path,
    *,
    first_col_title: str,
    row_labels: Sequence,
    label_format: str | None,
    vector,
    channel_names: Mapping[int, str],
    states: Sequence[str],
    table_prefix: str,
    decimals: int,
    pad_to_equal_width: bool,
    sex_map: Mapping[str, str] | None,
) -> Path:
    """Write one sheet per channel x state: a label column, then a block of columns per group."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if "time" in df.columns and df["time"].nunique() > 1:
        raise ValueError("Filter the table to one time point first, e.g. df[df['time'] == 'pre'].")

    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    palette = ["D9E8FB", "FBE0D9", "DFF5D8", "F5E8D8", "EAD8F5", "F5F0D8"]
    header_font = Font(name="Arial", bold=True, size=11)
    group_font = Font(name="Arial", bold=True, size=12)
    data_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center")
    num_fmt = "0." + "0" * decimals

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for ch, ch_name in channel_names.items():
        for state in states:
            sub = df[(df["ch"] == ch) & (df["state"] == state)]
            name = f"{table_prefix}_{ch_name}_{state}" if table_prefix else f"{ch_name}_{state}"
            ws = wb.create_sheet(name[:31])

            corner = ws.cell(row=2, column=1, value=first_col_title)
            corner.font, corner.alignment, corner.border = header_font, center, border

            by_group = {g: _sort_animals(sub[sub["group"] == g], sex_map) for g in groups}
            widest = max((len(r) for r in by_group.values()), default=0)

            col = 2
            blocks = []
            for gi, group in enumerate(groups):
                rows = by_group[group]
                n = widest if pad_to_equal_width else len(rows)
                if n == 0:
                    continue
                fill = PatternFill("solid", fgColor=palette[gi % len(palette)])
                ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + n - 1)
                gcell = ws.cell(row=1, column=col, value=group)
                gcell.font, gcell.alignment = group_font, center
                vectors = []
                for j in range(n):
                    present = j < len(rows)
                    head = ws.cell(
                        row=2, column=col + j, value=rows.iloc[j]["animal"] if present else None
                    )
                    head.font, head.fill, head.alignment, head.border = (
                        header_font,
                        fill,
                        center,
                        border,
                    )
                    ws.cell(row=1, column=col + j).fill = fill
                    ws.cell(row=1, column=col + j).border = border
                    vectors.append(vector(rows.iloc[j]) if present else None)
                blocks.append((col, vectors))
                col += n

            for r, label in enumerate(row_labels, start=3):
                cell = ws.cell(row=r, column=1, value=float(label) if label_format else label)
                cell.font = data_font
                if label_format:
                    cell.number_format = label_format
            for start, vectors in blocks:
                for j, values in enumerate(vectors):
                    if values is None:
                        continue  # padding column stays blank
                    for r, value in enumerate(values, start=3):
                        cell = ws.cell(row=r, column=start + j, value=float(value))
                        cell.font, cell.number_format = data_font, num_fmt

            ws.freeze_panes = "B3"
            ws.column_dimensions["A"].width = 14
            for c in range(2, col):
                ws.column_dimensions[get_column_letter(c)].width = 12

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def export_spectra_prism(
    df: pd.DataFrame,
    value_col: str,
    groups: Sequence[str],
    output_path: str | Path,
    *,
    channel_names: Mapping[int, str] = CHANNEL_NAMES,
    states: Sequence[str] = STATES,
    table_prefix: str = "",
    decimals: int = 6,
    pad_to_equal_width: bool = True,
    sex_map: Mapping[str, str] | None = None,
) -> Path:
    """Write full spectra to an .xlsx laid out for a Prism XY table with replicates.

    One sheet per channel x state: a frequency column, then a block of columns per group (in the
    order of ``groups``), one column per animal. ``df`` must hold a single time point.
    """
    freqs = get_freqs(df)
    return _write_prism(
        df,
        list(groups),
        output_path,
        first_col_title="Frequency (Hz)",
        row_labels=list(freqs),
        label_format="0.0",
        vector=lambda row: row[value_col],
        channel_names=channel_names,
        states=states,
        table_prefix=table_prefix,
        decimals=decimals,
        pad_to_equal_width=pad_to_equal_width,
        sex_map=sex_map,
    )


def export_band_power_prism(
    table: pd.DataFrame,
    bands: Sequence[str],
    groups: Sequence[str],
    output_path: str | Path,
    *,
    channel_names: Mapping[int, str] = CHANNEL_NAMES,
    states: Sequence[str] = STATES,
    table_prefix: str = "",
    decimals: int = 6,
    pad_to_equal_width: bool = True,
    sex_map: Mapping[str, str] | None = None,
) -> Path:
    """Write band power to an .xlsx laid out for a Prism grouped table.

    One sheet per channel x state: a row per band, then a block of columns per group, one column
    per animal. ``table`` comes from :func:`band_power_table` and must hold a single time point.
    """
    bands = list(bands)
    return _write_prism(
        table,
        list(groups),
        output_path,
        first_col_title="Band",
        row_labels=bands,
        label_format=None,
        vector=lambda row: [row[b] for b in bands],
        channel_names=channel_names,
        states=states,
        table_prefix=table_prefix,
        decimals=decimals,
        pad_to_equal_width=pad_to_equal_width,
        sex_map=sex_map,
    )
