"""Tests for ``notebooks/psd_utils.py`` (helpers for the PSD analysis notebook)."""

from __future__ import annotations

import pickle
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "notebooks"))
import psd_utils as u  # noqa: E402

FREQS = np.arange(0, 125.5, 0.5)


def _record(animal: str, state: str, scale: float = 1.0, **extra) -> dict:
    psd = np.stack([np.full(FREQS.size, scale * (ch + 1), dtype=float) for ch in range(3)])
    rec = {
        "path": Path(f"/data/{animal}_P001_PREPROCESSED.edf"),
        "behavior": ("threshold", state),
        "freqs": FREQS,
        "psd": psd,
        "n_estimates": 100,
        "spindle_out_frac": 0.1,
        "comm_out_frac": 0.0,
        "sd_out_frac": 0.0,
        "swd_out_frac": 0.0,
        "ann_out_frac": 0.0,
    }
    rec.update(extra)
    return rec


def _records(animals=("A1", "A2", "B1", "B2")) -> list[dict]:
    return [_record(a, s, scale=1 + i) for i, a in enumerate(animals) for s in u.STATES]


def test_records_to_df_has_one_row_per_animal_state_channel():
    df = u.records_to_df(_records(), apply_inverse_filter=False)
    assert len(df) == 4 * 3 * 3
    assert set(df["channel"]) == {"LFC", "RSC", "LSC"}
    assert set(df["animal"]) == {"A1", "A2", "B1", "B2"}
    assert np.array_equal(u.get_freqs(df), FREQS)


def test_inverse_filter_scales_only_0p5_to_3p5_hz_and_leaves_records_alone():
    recs = _records(("A1",))
    df = u.records_to_df(recs, apply_inverse_filter=True)
    spec = df.iloc[0]["spectra"]
    inside = (FREQS >= 0.5) & (FREQS <= 3.5)
    assert np.allclose(spec[inside], 1.0 * u.INVERSE_FILTER)
    assert np.allclose(spec[~inside], 1.0)
    assert np.allclose(recs[0]["psd"][0], 1.0)  # the input is not modified
    again = u.records_to_df(recs, apply_inverse_filter=True)
    assert np.allclose(again.iloc[0]["spectra"], spec)  # and calling twice gives the same answer


def test_inverse_filter_needs_half_hertz_bins():
    coarse = np.arange(0, 125.1, 1.0)
    rec = _record("A1", "wake")
    rec["freqs"] = coarse
    rec["psd"] = np.ones((3, coarse.size))
    with pytest.raises(ValueError, match="0.5 Hz"):
        u.records_to_df([rec])
    assert len(u.records_to_df([rec], apply_inverse_filter=False)) == 3


def test_excluded_animals_are_left_out_unless_asked_and_listed_with_their_reasons(capsys):
    recs = _records(("A1", "A2", "A3"))
    for r in recs:
        if "A2" in str(r["path"]):
            r["excluded"], r["exclusion_reason"] = True, "Sleep QC: bad file"
        if "A3" in str(r["path"]):
            r["excluded"], r["exclusion_reason"] = True, "PSD review: extra bad"
    df = u.records_to_df(recs, apply_inverse_filter=False)
    assert set(df["animal"]) == {"A1"}
    out = capsys.readouterr().out
    assert "A2: Sleep QC: bad file" in out and "A3: PSD review: extra bad" in out
    kept = u.records_to_df(recs, apply_inverse_filter=False, include_excluded=True)
    assert set(kept["animal"]) == {"A1", "A2", "A3"}
    assert kept.loc[kept["animal"] == "A3", "excluded"].all()
    assert kept.loc[kept["animal"] == "A3", "exclusion_reason"].iloc[0] == "PSD review: extra bad"


def test_load_psd_runs_adds_time_and_keeps_the_frequency_axis(tmp_path: Path):
    for label in ("pre", "post"):
        with open(tmp_path / f"{label}.pkl", "wb") as f:
            pickle.dump(_records(("A1", "B1")), f)
    df = u.load_psd_runs(
        {"pre": tmp_path / "pre.pkl", "post": tmp_path / "post.pkl"}, apply_inverse_filter=False
    )
    assert set(df["time"]) == {"pre", "post"}
    assert np.array_equal(u.get_freqs(df), FREQS)
    assert np.array_equal(u.get_freqs(df[df["time"] == "pre"]), FREQS)


def test_assign_groups_labels_everyone_and_reports_the_rest(capsys):
    df = u.records_to_df(_records(), apply_inverse_filter=False)
    out = u.assign_groups(df, {"WT": ["A1", "A2"], "MUT": ["B1", "ZZ9"]})
    assert out.loc[out["animal"] == "A1", "group"].iloc[0] == "WT"
    assert out.loc[out["animal"] == "B2", "group"].iloc[0] == "unassigned"
    text = capsys.readouterr().out
    assert "B2" in text and "ZZ9" in text
    assert np.array_equal(u.get_freqs(out), FREQS)


def test_interpolation_bridges_the_line_noise_peak():
    rec = _record("A1", "wake")
    spike = np.ones(FREQS.size)
    spike[(FREQS >= 59) & (FREQS <= 61)] = 50
    rec["psd"] = np.stack([spike] * 3)
    df = u.records_to_df([rec], apply_inverse_filter=False)
    out = u.interpolate_line_noise(df, ranges=[(58, 62)])
    i60 = int(np.argmin(np.abs(FREQS - 60)))
    assert out.iloc[0]["spectra"][i60] == 50
    assert out.iloc[0]["interp_spectra"][i60] < 5


def test_default_interpolation_bridges_56_to_64_hz():
    spike = np.ones(FREQS.size)
    spike[(FREQS >= 56) & (FREQS <= 64)] = 50
    rec = _record("A1", "wake")
    rec["psd"] = np.stack([spike] * 3)
    df = u.records_to_df([rec], apply_inverse_filter=False)
    default = u.interpolate_line_noise(df)
    narrow = u.interpolate_line_noise(df, ranges=[(58, 62)])
    i57 = int(np.argmin(np.abs(FREQS - 57)))
    assert default.iloc[0]["interp_spectra"][i57] < 5
    assert narrow.iloc[0]["interp_spectra"][i57] == 50


def test_normalized_spectra_have_unit_power_in_the_band():
    rec = _record("A1", "wake", scale=7.0)
    df = u.records_to_df([rec], apply_inverse_filter=False)
    df = u.interpolate_line_noise(df)
    df = u.normalize_spectra(df, band=(1, 100))
    freqs = u.get_freqs(df)
    for spec in df["norm_spectra"]:
        assert u.compute_bandpower(spec, freqs, (1, 100)) == pytest.approx(1.0)


def test_band_power_is_the_area_under_the_curve():
    freqs = FREQS
    flat = np.full(freqs.size, 2.0)
    assert u.compute_bandpower(flat, freqs, (4, 8)) == pytest.approx(8.0)  # 2 uV^2/Hz over 4 Hz


def test_band_power_table_has_a_column_per_band():
    df = u.records_to_df(_records(), apply_inverse_filter=False)
    df = u.assign_groups(df, {"WT": ["A1", "A2"], "MUT": ["B1", "B2"]})
    table = u.band_power_table(df, "spectra")
    assert list(u.DEFAULT_BANDS) == [c for c in table.columns if c in u.DEFAULT_BANDS]
    assert len(table) == len(df)
    assert {"animal", "state", "group", "ch"} <= set(table.columns)


def _grouped_df(animals=("A1", "A2", "B1", "B2", "B3")):
    df = u.records_to_df(_records(animals), apply_inverse_filter=False)
    groups = {
        "WT": [a for a in animals if a.startswith("A")],
        "MUT": [a for a in animals if a.startswith("B")],
    }
    df = u.assign_groups(df, groups)
    df = u.interpolate_line_noise(df)
    return u.normalize_spectra(df)


def test_spectra_plots_return_figures(tmp_path: Path):
    df = _grouped_df()
    fig = u.plot_spectra(
        df, color_map={"WT": "tab:green", "MUT": "tab:blue"}, save_path=tmp_path / "s.png"
    )
    assert (tmp_path / "s.png").exists()
    assert len(fig.axes) >= 9
    u.plot_spectra(df, show_individuals=True, use_median=True, log_y=False)
    u.plot_spectra_by_animal(df, column="norm_spectra")


def test_boxplot_band_power_labels_an_outlier(tmp_path: Path):
    animals = tuple(f"A{i}" for i in range(8))
    recs = [_record(a, s, scale=1.0) for a in animals for s in u.STATES]
    recs[0]["psd"] = recs[0]["psd"] * 50  # one animal far above the rest
    df = u.records_to_df(recs, apply_inverse_filter=False)
    df = u.assign_groups(df, {"WT": list(animals)})
    table = u.band_power_table(df, "spectra")
    fig, axs = u.boxplot_band_power(
        table, "4-8", "Power", "t", x="group", save_path=tmp_path / "b.png"
    )
    texts = [t.get_text() for ax in axs.flat for t in ax.texts]
    assert "A0" in texts
    assert (tmp_path / "b.png").exists()


def test_prism_exports_have_a_sheet_per_channel_and_state(tmp_path: Path):
    import openpyxl

    df = _grouped_df()
    spectra = u.export_spectra_prism(df, "interp_spectra", ["WT", "MUT"], tmp_path / "sp.xlsx")
    wb = openpyxl.load_workbook(spectra)
    assert set(wb.sheetnames) == {f"{c}_{s}" for c in ("LFC", "RSC", "LSC") for s in u.STATES}
    ws = wb["LFC_wake"]
    assert ws["A2"].value == "Frequency (Hz)"
    assert ws["B1"].value == "WT"
    header = [c.value for c in ws[2]]
    assert header[1:3] == ["A1", "A2"]  # WT block: alphabetical
    assert ws.max_row == 2 + FREQS.size

    table = u.band_power_table(df, "interp_spectra")
    bands = u.export_band_power_prism(
        table, list(u.DEFAULT_BANDS), ["WT", "MUT"], tmp_path / "bp.xlsx"
    )
    ws = openpyxl.load_workbook(bands)["RSC_rem"]
    assert ws["A2"].value == "Band"
    assert [ws.cell(row=r, column=1).value for r in range(3, 9)] == list(u.DEFAULT_BANDS)


def test_prism_export_refuses_a_table_with_several_time_points(tmp_path: Path):
    df = _grouped_df()
    df2 = pd.concat([df.assign(time="pre"), df.assign(time="post")], ignore_index=True)
    df2.attrs = dict(df.attrs)
    with pytest.raises(ValueError, match="one time point"):
        u.export_spectra_prism(df2, "interp_spectra", ["WT", "MUT"], tmp_path / "x.xlsx")
